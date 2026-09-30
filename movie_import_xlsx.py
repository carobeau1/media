"""Read movie IDs from column A of the first sheet of an uploaded workbook."""
from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook

MAX_IDS = 10000
MAX_FILE_BYTES = 5 * 1024 * 1024


def read_ids(upload):
    if not upload or not upload.filename or not upload.filename.lower().endswith('.xlsx'):
        raise ValueError('Select an .xlsx spreadsheet.')
    data = upload.stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('The spreadsheet must be 5 MB or smaller.')
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            ids = []
            for number, (value,) in enumerate(sheet.iter_rows(min_col=1, max_col=1, values_only=True), 1):
                if value is None or isinstance(value, str) and not value.strip():
                    continue
                if isinstance(value, bool):
                    movie_id = ''
                elif isinstance(value, int):
                    movie_id = str(value)
                elif isinstance(value, float) and value.is_integer():
                    movie_id = str(int(value))
                else:
                    movie_id = str(value).strip()
                if not movie_id.isascii() or not movie_id.isdigit() or len(movie_id) > 15:
                    raise ValueError(f'Invalid movie ID in cell A{number}.')
                ids.append(movie_id)
                if len(ids) > MAX_IDS:
                    raise ValueError('Import up to 10,000 movie IDs at a time.')
            if not ids:
                raise ValueError('Column A of the first sheet has no movie IDs.')
            return ids
        finally:
            workbook.close()
    except (BadZipFile, KeyError, IndexError, OSError) as exc:
        raise ValueError('Could not read the .xlsx file. Check that it is a valid spreadsheet.') from exc
